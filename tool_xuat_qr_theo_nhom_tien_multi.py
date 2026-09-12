# -*- coding: utf-8 -*-
"""
TOOL XUẤT QR THEO NHÓM TIỀN - MULTI FILE
=========================================
Chức năng:
- Import 1 hoặc nhiều file Excel cùng lúc
- Có thể thêm tiếp file hoặc xóa file khỏi danh sách
- Tự quét các mức tiền có trong tất cả file
- Chọn nhóm tiền cần xuất
- Xuất QR theo thư mục nhóm tiền
- Có tùy chọn gom chung hoặc chia thêm theo tên file nguồn
- Tự tránh trùng tên ảnh
- Tạo file thống kê
- Không sửa file Excel gốc

Yêu cầu:
    pip install openpyxl

Mặc định cấu trúc Excel:
    Sheet: DanhSach
    Cột B: Họ tên
    Cột D: CCCD
    Cột F: Số tiền
"""

import os
import re
import shutil
import zipfile
import traceback
import xml.etree.ElementTree as ET
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import load_workbook


APP_TITLE = "XUẤT QR THEO NHÓM TIỀN"

DEFAULT_SHEET = "DanhSach"
DEFAULT_COL_NAME = "B"
DEFAULT_COL_CCCD = "D"
DEFAULT_COL_MONEY = "F"


# ============================================================
# TIỆN ÍCH DỮ LIỆU
# ============================================================

def col_letter_to_number(col):
    col = str(col).strip().upper()
    n = 0
    for ch in col:
        if not ("A" <= ch <= "Z"):
            raise ValueError(f"Cột không hợp lệ: {col}")
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n


def safe_filename(text):
    text = str(text).strip()
    text = re.sub(r'[\\/:*?"<>|]', "_", text)
    text = re.sub(r"\s+", "_", text)
    return text.strip("._ ")


def normalize_name(name):
    return safe_filename(str(name).strip().upper())


def normalize_cccd(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = str(value).strip()

    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]

    return text


def normalize_money(value):
    """
    Các dạng hỗ trợ:
    350000
    350000.0
    350.000
    350,000
    350.000 đ
    350000 VND
    """
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        return int(round(value))

    text = str(value).strip()
    if not text:
        return None

    text = re.sub(r"(?i)vnd", "", text)
    text = text.replace("₫", "").replace("đ", "").replace("Đ", "")
    text = text.replace(" ", "")

    # 350.000 / 1.150.000 / 350,000
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", text):
        return int(re.sub(r"[.,]", "", text))

    # 350000.0 / 350000,0
    if re.fullmatch(r"\d+[.,]0+", text):
        return int(re.split(r"[.,]", text)[0])

    if text.isdigit():
        return int(text)

    # fallback: lấy chữ số
    digits = re.sub(r"\D", "", text)
    return int(digits) if digits else None


def format_money_vn(amount):
    return f"{int(amount):,}".replace(",", ".")


# ============================================================
# ĐỌC IMAGE TỪ XLSX
# ============================================================

def get_sheet_and_drawing_map(xlsx_path):
    ns_main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_rel_doc = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ns_rel_pkg = "http://schemas.openxmlformats.org/package/2006/relationships"

    result = {}

    with zipfile.ZipFile(xlsx_path, "r") as z:
        workbook_xml = ET.fromstring(z.read("xl/workbook.xml"))
        workbook_rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))

        rel_map = {}
        for rel in workbook_rels.findall("{%s}Relationship" % ns_rel_pkg):
            rid = rel.attrib["Id"]
            target = rel.attrib["Target"]

            if target.startswith("/"):
                target = target.lstrip("/")
            elif not target.startswith("xl/"):
                target = "xl/" + target

            rel_map[rid] = target

        sheets = workbook_xml.find("{%s}sheets" % ns_main)

        for sh in sheets.findall("{%s}sheet" % ns_main):
            sheet_name = sh.attrib["name"]
            rid = sh.attrib["{%s}id" % ns_rel_doc]

            if rid not in rel_map:
                continue

            sheet_path = rel_map[rid]
            rels_path = (
                os.path.dirname(sheet_path).replace("\\", "/")
                + "/_rels/"
                + os.path.basename(sheet_path)
                + ".rels"
            )

            drawing_path = None

            if rels_path in z.namelist():
                sheet_xml = ET.fromstring(z.read(sheet_path))
                drawing_el = sheet_xml.find("{%s}drawing" % ns_main)

                if drawing_el is not None:
                    drawing_rid = drawing_el.attrib.get("{%s}id" % ns_rel_doc)
                    rels_xml = ET.fromstring(z.read(rels_path))

                    for rel in rels_xml.findall("{%s}Relationship" % ns_rel_pkg):
                        if rel.attrib.get("Id") == drawing_rid:
                            target = rel.attrib.get("Target", "")
                            base_dir = os.path.dirname(sheet_path).replace("\\", "/")
                            drawing_path = os.path.normpath(
                                os.path.join(base_dir, target)
                            ).replace("\\", "/")
                            break

            result[sheet_name] = {
                "sheet_path": sheet_path,
                "drawing_path": drawing_path,
            }

    return result


def get_image_relations(zip_file, drawing_xml_path):
    ns_rel_pkg = "http://schemas.openxmlformats.org/package/2006/relationships"

    drawing_dir = os.path.dirname(drawing_xml_path).replace("\\", "/")
    drawing_name = os.path.basename(drawing_xml_path)
    rels_path = drawing_dir + "/_rels/" + drawing_name + ".rels"

    rel_map = {}
    if rels_path not in zip_file.namelist():
        return rel_map

    rels_xml = ET.fromstring(zip_file.read(rels_path))

    for rel in rels_xml.findall("{%s}Relationship" % ns_rel_pkg):
        rid = rel.attrib["Id"]
        target = rel.attrib["Target"]
        image_path = os.path.normpath(
            os.path.join(drawing_dir, target)
        ).replace("\\", "/")
        rel_map[rid] = image_path

    return rel_map


def extract_images_by_row(xlsx_path, sheet_name, temp_folder):
    ns = {
        "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }

    mapping = get_sheet_and_drawing_map(xlsx_path)

    if sheet_name not in mapping:
        raise Exception(f"Không tìm thấy sheet '{sheet_name}'")

    drawing_xml_path = mapping[sheet_name]["drawing_path"]

    if not drawing_xml_path:
        raise Exception(f"Sheet '{sheet_name}' không có ảnh.")

    results = []

    with zipfile.ZipFile(xlsx_path, "r") as z:
        drawing_xml = ET.fromstring(z.read(drawing_xml_path))
        rel_map = get_image_relations(z, drawing_xml_path)

        anchors = (
            drawing_xml.findall("xdr:oneCellAnchor", ns)
            + drawing_xml.findall("xdr:twoCellAnchor", ns)
        )

        Path(temp_folder).mkdir(parents=True, exist_ok=True)

        stt = 1
        for anchor in anchors:
            from_el = anchor.find("xdr:from", ns)
            if from_el is None:
                continue

            row_el = from_el.find("xdr:row", ns)
            if row_el is None:
                continue

            row_excel = int(row_el.text) + 1

            pic = anchor.find("xdr:pic", ns)
            if pic is None:
                continue

            blip = pic.find(".//a:blip", ns)
            if blip is None:
                continue

            embed = blip.attrib.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
            )

            if not embed or embed not in rel_map:
                continue

            image_in_zip = rel_map[embed]
            if image_in_zip not in z.namelist():
                continue

            ext = Path(image_in_zip).suffix.lower() or ".png"
            temp_path = Path(temp_folder) / f"img_{stt:05d}{ext}"

            with z.open(image_in_zip) as src, open(temp_path, "wb") as dst:
                shutil.copyfileobj(src, dst)

            results.append({
                "row": row_excel,
                "temp_path": temp_path,
            })

            stt += 1

    return results


# ============================================================
# GUI
# ============================================================

class QRGroupTool:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("980x760")
        self.root.minsize(900, 700)

        self.files = []
        self.detected_amounts = {}

        self.sheet_var = tk.StringVar(value=DEFAULT_SHEET)
        self.name_col_var = tk.StringVar(value=DEFAULT_COL_NAME)
        self.cccd_col_var = tk.StringVar(value=DEFAULT_COL_CCCD)
        self.money_col_var = tk.StringVar(value=DEFAULT_COL_MONEY)
        self.output_var = tk.StringVar()
        self.subfolder_source_var = tk.BooleanVar(value=False)
        self.summary_var = tk.BooleanVar(value=True)
        self.status_var = tk.StringVar(value="Chưa import file Excel.")

        self.build_ui()

    def build_ui(self):
        pad = 12

        ttk.Label(
            self.root,
            text=APP_TITLE,
            font=("Segoe UI", 18, "bold"),
        ).pack(pady=(15, 8))

        # FILES
        f1 = ttk.LabelFrame(self.root, text="1. Import file Excel")
        f1.pack(fill="both", padx=pad, pady=6)

        toolbar = ttk.Frame(f1)
        toolbar.pack(fill="x", padx=8, pady=8)

        ttk.Button(
            toolbar,
            text="Import nhiều file...",
            command=self.add_files
        ).pack(side="left", padx=(0, 6))

        ttk.Button(
            toolbar,
            text="Import cả thư mục...",
            command=self.add_folder
        ).pack(side="left", padx=6)

        ttk.Button(
            toolbar,
            text="Xóa file đã chọn",
            command=self.remove_selected
        ).pack(side="left", padx=6)

        ttk.Button(
            toolbar,
            text="Xóa tất cả",
            command=self.clear_files
        ).pack(side="left", padx=6)

        self.file_tree = ttk.Treeview(
            f1,
            columns=("name", "path"),
            show="headings",
            height=8
        )
        self.file_tree.heading("name", text="Tên file")
        self.file_tree.heading("path", text="Đường dẫn")
        self.file_tree.column("name", width=300)
        self.file_tree.column("path", width=600)
        self.file_tree.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        # CONFIG
        f2 = ttk.LabelFrame(self.root, text="2. Cấu trúc Excel")
        f2.pack(fill="x", padx=pad, pady=6)

        ttk.Label(f2, text="Tên sheet:").grid(row=0, column=0, padx=8, pady=8, sticky="w")
        ttk.Entry(f2, textvariable=self.sheet_var, width=20).grid(row=0, column=1, pady=8)

        ttk.Label(f2, text="Cột Họ tên:").grid(row=0, column=2, padx=(20, 8), pady=8)
        ttk.Entry(f2, textvariable=self.name_col_var, width=8).grid(row=0, column=3, pady=8)

        ttk.Label(f2, text="Cột CCCD:").grid(row=0, column=4, padx=(20, 8), pady=8)
        ttk.Entry(f2, textvariable=self.cccd_col_var, width=8).grid(row=0, column=5, pady=8)

        ttk.Label(f2, text="Cột Số tiền:").grid(row=0, column=6, padx=(20, 8), pady=8)
        ttk.Entry(f2, textvariable=self.money_col_var, width=8).grid(row=0, column=7, pady=8)

        ttk.Button(
            f2,
            text="Quét nhóm tiền",
            command=self.scan_amounts
        ).grid(row=0, column=8, padx=12, pady=8)

        # AMOUNTS
        f3 = ttk.LabelFrame(self.root, text="3. Chọn nhóm tiền cần xuất")
        f3.pack(fill="both", expand=True, padx=pad, pady=6)

        top3 = ttk.Frame(f3)
        top3.pack(fill="x", padx=8, pady=6)

        ttk.Button(
            top3,
            text="Chọn tất cả",
            command=lambda: self.set_all_amounts(True)
        ).pack(side="left", padx=(0, 6))

        ttk.Button(
            top3,
            text="Bỏ chọn tất cả",
            command=lambda: self.set_all_amounts(False)
        ).pack(side="left", padx=6)

        self.amount_tree = ttk.Treeview(
            f3,
            columns=("selected", "amount", "count"),
            show="headings",
            height=8
        )
        self.amount_tree.heading("selected", text="Xuất?")
        self.amount_tree.heading("amount", text="Mức tiền")
        self.amount_tree.heading("count", text="Số dòng")
        self.amount_tree.column("selected", width=80, anchor="center")
        self.amount_tree.column("amount", width=220, anchor="center")
        self.amount_tree.column("count", width=120, anchor="center")
        self.amount_tree.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self.amount_tree.bind("<Double-1>", self.toggle_amount)

        # OUTPUT
        f4 = ttk.LabelFrame(self.root, text="4. Xuất kết quả")
        f4.pack(fill="x", padx=pad, pady=6)

        ttk.Label(f4, text="Thư mục output:").grid(
            row=0, column=0, padx=8, pady=(8, 4), sticky="w"
        )
        ttk.Entry(f4, textvariable=self.output_var, width=75).grid(
            row=1, column=0, padx=(8, 6), pady=(0, 8), sticky="we"
        )
        ttk.Button(
            f4, text="Chọn...",
            command=self.choose_output
        ).grid(row=1, column=1, padx=(0, 8), pady=(0, 8))

        ttk.Checkbutton(
            f4,
            text="Chia thêm thư mục theo file nguồn",
            variable=self.subfolder_source_var
        ).grid(row=2, column=0, padx=8, pady=4, sticky="w")

        ttk.Checkbutton(
            f4,
            text="Tạo file _THONG_KE.txt",
            variable=self.summary_var
        ).grid(row=3, column=0, padx=8, pady=(4, 8), sticky="w")

        f4.columnconfigure(0, weight=1)

        ttk.Button(
            self.root,
            text="XUẤT QR",
            command=self.export_qr
        ).pack(pady=(10, 5), ipadx=32, ipady=8)

        ttk.Label(
            self.root,
            textvariable=self.status_var,
            wraplength=920,
            justify="center"
        ).pack(pady=(4, 12))

    # --------------------------------------------------------

    def refresh_file_tree(self):
        for item in self.file_tree.get_children():
            self.file_tree.delete(item)

        for p in self.files:
            self.file_tree.insert("", "end", values=(p.name, str(p)))

        self.status_var.set(f"Đã import {len(self.files)} file Excel.")

    def add_files(self):
        paths = filedialog.askopenfilenames(
            title="Chọn một hoặc nhiều file Excel",
            filetypes=[
                ("Excel files", "*.xlsx *.xlsm"),
                ("XLSX", "*.xlsx"),
                ("XLSM", "*.xlsm"),
            ]
        )
        if not paths:
            return

        current = {str(p.resolve()).lower() for p in self.files}

        for x in paths:
            p = Path(x)
            key = str(p.resolve()).lower()
            if key not in current:
                self.files.append(p)
                current.add(key)

        self.files.sort(key=lambda p: p.name.lower())
        self.refresh_file_tree()

        if not self.output_var.get() and self.files:
            self.output_var.set(str(self.files[0].parent / "output_qr"))

    def add_folder(self):
        folder = filedialog.askdirectory(title="Chọn thư mục chứa Excel")
        if not folder:
            return

        folder = Path(folder)
        found = list(folder.glob("*.xlsx")) + list(folder.glob("*.xlsm"))

        current = {str(p.resolve()).lower() for p in self.files}
        for p in found:
            if p.name.startswith("~$"):
                continue
            key = str(p.resolve()).lower()
            if key not in current:
                self.files.append(p)
                current.add(key)

        self.files.sort(key=lambda p: p.name.lower())
        self.refresh_file_tree()

        if not self.output_var.get():
            self.output_var.set(str(folder / "output_qr"))

    def remove_selected(self):
        selected = self.file_tree.selection()
        if not selected:
            return

        remove_paths = {
            self.file_tree.item(i)["values"][1]
            for i in selected
        }

        self.files = [
            p for p in self.files
            if str(p) not in remove_paths
        ]
        self.refresh_file_tree()

    def clear_files(self):
        self.files = []
        self.detected_amounts = {}
        self.refresh_file_tree()
        self.refresh_amount_tree()

    def choose_output(self):
        folder = filedialog.askdirectory(title="Chọn thư mục output")
        if folder:
            self.output_var.set(folder)

    # --------------------------------------------------------

    def get_columns(self):
        return (
            col_letter_to_number(self.name_col_var.get()),
            col_letter_to_number(self.cccd_col_var.get()),
            col_letter_to_number(self.money_col_var.get()),
        )

    def scan_amounts(self):
        if not self.files:
            messagebox.showwarning("Chưa có file", "Hãy import file Excel trước.")
            return

        try:
            _, _, money_col = self.get_columns()
            sheet_name = self.sheet_var.get().strip()
            if not sheet_name:
                raise ValueError("Chưa nhập tên sheet.")

            counts = {}
            errors = []

            for idx, path in enumerate(self.files, start=1):
                self.status_var.set(
                    f"Đang quét {idx}/{len(self.files)}: {path.name}"
                )
                self.root.update_idletasks()

                try:
                    wb = load_workbook(path, data_only=True, read_only=True)

                    if sheet_name not in wb.sheetnames:
                        errors.append(f"{path.name}: không có sheet '{sheet_name}'")
                        wb.close()
                        continue

                    ws = wb[sheet_name]

                    for row in range(2, ws.max_row + 1):
                        amount = normalize_money(
                            ws.cell(row=row, column=money_col).value
                        )
                        if amount is not None:
                            counts[amount] = counts.get(amount, 0) + 1

                    wb.close()

                except Exception as e:
                    errors.append(f"{path.name}: {e}")

            self.detected_amounts = {
                amount: {
                    "count": count,
                    "selected": True,
                }
                for amount, count in sorted(counts.items(), reverse=True)
            }

            self.refresh_amount_tree()

            msg = (
                f"Đã quét {len(self.files)} file. "
                f"Tìm thấy {len(self.detected_amounts)} nhóm tiền."
            )

            if errors:
                msg += f" Có {len(errors)} file lỗi."

            self.status_var.set(msg)

            if errors:
                messagebox.showwarning(
                    "Quét hoàn tất nhưng có lỗi",
                    msg + "\n\n" + "\n".join(errors[:15])
                )

        except Exception as e:
            messagebox.showerror("Lỗi", str(e))

    def refresh_amount_tree(self):
        for item in self.amount_tree.get_children():
            self.amount_tree.delete(item)

        for amount, info in self.detected_amounts.items():
            self.amount_tree.insert(
                "",
                "end",
                iid=str(amount),
                values=(
                    "✓" if info["selected"] else "",
                    format_money_vn(amount),
                    info["count"],
                )
            )

    def toggle_amount(self, event=None):
        sel = self.amount_tree.selection()
        if not sel:
            return

        for iid in sel:
            amount = int(iid)
            self.detected_amounts[amount]["selected"] = (
                not self.detected_amounts[amount]["selected"]
            )

        self.refresh_amount_tree()

    def set_all_amounts(self, value):
        for info in self.detected_amounts.values():
            info["selected"] = value
        self.refresh_amount_tree()

    # --------------------------------------------------------

    def unique_output_path(self, folder, filename):
        folder = Path(folder)
        path = folder / filename

        if not path.exists():
            return path

        stem = path.stem
        suffix = path.suffix

        n = 2
        while True:
            candidate = folder / f"{stem}_{n}{suffix}"
            if not candidate.exists():
                return candidate
            n += 1

    def export_qr(self):
        if not self.files:
            messagebox.showwarning("Chưa có file", "Hãy import file Excel.")
            return

        if not self.detected_amounts:
            messagebox.showwarning(
                "Chưa quét nhóm tiền",
                "Bấm 'Quét nhóm tiền' trước."
            )
            return

        selected_amounts = {
            a for a, info in self.detected_amounts.items()
            if info["selected"]
        }

        if not selected_amounts:
            messagebox.showwarning(
                "Chưa chọn nhóm",
                "Hãy chọn ít nhất một nhóm tiền."
            )
            return

        output_root = self.output_var.get().strip()
        if not output_root:
            messagebox.showwarning(
                "Chưa có output",
                "Hãy chọn thư mục output."
            )
            return

        try:
            name_col, cccd_col, money_col = self.get_columns()
            sheet_name = self.sheet_var.get().strip()

            output_root = Path(output_root)
            output_root.mkdir(parents=True, exist_ok=True)

            temp_root = output_root / "__temp__"
            if temp_root.exists():
                shutil.rmtree(temp_root, ignore_errors=True)
            temp_root.mkdir(parents=True, exist_ok=True)

            stats = {}
            skipped = []
            errors = []
            total_exported = 0

            for file_index, excel_path in enumerate(self.files, start=1):
                self.status_var.set(
                    f"Đang xuất {file_index}/{len(self.files)}: {excel_path.name}"
                )
                self.root.update_idletasks()

                file_temp = temp_root / f"file_{file_index:04d}"

                try:
                    wb = load_workbook(
                        excel_path,
                        data_only=True,
                        read_only=False
                    )

                    if sheet_name not in wb.sheetnames:
                        errors.append(
                            f"{excel_path.name}: không có sheet '{sheet_name}'"
                        )
                        wb.close()
                        continue

                    ws = wb[sheet_name]

                    images = extract_images_by_row(
                        str(excel_path),
                        sheet_name,
                        str(file_temp)
                    )

                    for item in images:
                        row = item["row"]
                        temp_path = Path(item["temp_path"])

                        if row <= 1:
                            temp_path.unlink(missing_ok=True)
                            continue

                        name = ws.cell(row=row, column=name_col).value
                        cccd = ws.cell(row=row, column=cccd_col).value
                        money_raw = ws.cell(row=row, column=money_col).value
                        amount = normalize_money(money_raw)

                        if amount not in selected_amounts:
                            temp_path.unlink(missing_ok=True)
                            continue

                        if not name:
                            temp_path.unlink(missing_ok=True)
                            skipped.append(
                                f"{excel_path.name} - dòng {row}: thiếu họ tên"
                            )
                            continue

                        if not cccd:
                            temp_path.unlink(missing_ok=True)
                            skipped.append(
                                f"{excel_path.name} - dòng {row}: thiếu CCCD"
                            )
                            continue

                        amount_folder = output_root / str(amount)

                        if self.subfolder_source_var.get():
                            source_folder = safe_filename(excel_path.stem)
                            final_folder = amount_folder / source_folder
                        else:
                            final_folder = amount_folder

                        final_folder.mkdir(parents=True, exist_ok=True)

                        ext = temp_path.suffix.lower() or ".png"
                        filename = safe_filename(
                            f"{normalize_name(name)}_{normalize_cccd(cccd)}"
                        ) + ext

                        final_path = self.unique_output_path(
                            final_folder,
                            filename
                        )

                        shutil.move(str(temp_path), str(final_path))

                        stats[amount] = stats.get(amount, 0) + 1
                        total_exported += 1

                    wb.close()

                except Exception as e:
                    errors.append(
                        f"{excel_path.name}: {e}"
                    )
                finally:
                    if file_temp.exists():
                        shutil.rmtree(file_temp, ignore_errors=True)

            shutil.rmtree(temp_root, ignore_errors=True)

            # Summary
            if self.summary_var.get():
                lines = []
                lines.append("THỐNG KÊ XUẤT QR THEO NHÓM TIỀN")
                lines.append("=" * 60)
                lines.append(f"Số file Excel: {len(self.files)}")
                lines.append("")

                for amount in sorted(stats.keys(), reverse=True):
                    lines.append(
                        f"{format_money_vn(amount)}: {stats[amount]} ảnh"
                    )

                lines.append("-" * 60)
                lines.append(f"TỔNG ẢNH: {total_exported}")
                lines.append(f"BỎ QUA: {len(skipped)}")
                lines.append(f"FILE LỖI: {len(errors)}")

                if skipped:
                    lines.append("")
                    lines.append("CHI TIẾT BỎ QUA:")
                    lines.extend("- " + x for x in skipped)

                if errors:
                    lines.append("")
                    lines.append("LỖI:")
                    lines.extend("- " + x for x in errors)

                (output_root / "_THONG_KE.txt").write_text(
                    "\n".join(lines),
                    encoding="utf-8-sig"
                )

            self.status_var.set(
                f"Hoàn thành: {total_exported} ảnh từ {len(self.files)} file Excel."
            )

            msg = (
                f"Đã xử lý {len(self.files)} file Excel.\n"
                f"Xuất thành công: {total_exported} ảnh QR.\n"
                f"Nhóm tiền đã xuất: {len(stats)}.\n"
                f"Output: {output_root}"
            )

            if skipped:
                msg += f"\nBỏ qua: {len(skipped)} dòng."
            if errors:
                msg += f"\nCó lỗi: {len(errors)} file."

            messagebox.showinfo("Hoàn thành", msg)

            try:
                os.startfile(str(output_root.resolve()))
            except Exception:
                pass

        except Exception as e:
            messagebox.showerror(
                "Lỗi",
                f"{e}\n\n{traceback.format_exc()[-2500:]}"
            )


def main():
    root = tk.Tk()

    try:
        style = ttk.Style()
        if "vista" in style.theme_names():
            style.theme_use("vista")
    except Exception:
        pass

    QRGroupTool(root)
    root.mainloop()


if __name__ == "__main__":
    main()
